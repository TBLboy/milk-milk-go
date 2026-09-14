import json
from io import BytesIO
import secrets
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.auth import require_admin
from app.db.models import Label, Material, PrintBatch, SystemSetting, User
from app.db.session import get_db
from app.services.audit import write_audit
from app.services.idempotency import execute_idempotent

router = APIRouter(prefix="/labels", tags=["labels"])
MATERIAL_EXCEL_HEADERS = ("material_code", "name_zh", "name_en", "shelf_life_months")


class PrintRequest(BaseModel):
    material_id: str = Field(min_length=1, max_length=32)
    quantity: int = Field(gt=0, le=10000)


def _now() -> datetime:
    return datetime.now(timezone.utc)


@router.post("/print-batches", status_code=201)
def create_print_batch(
    body: PrintRequest,
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> dict:
    def operation() -> dict:
        material = db.scalar(select(Material).where(Material.material_id == body.material_id))
        if material is None:
            raise HTTPException(status_code=422, detail={"code": "MATERIAL_NOT_FOUND", "message": "辅料不存在"})
        label_size = db.scalar(
            select(SystemSetting.value).where(SystemSetting.key == "label_size_mm")
        ) or "60x40"
        printed_at = _now()
        batch_id = f"PB-{printed_at.strftime('%Y%m%d%H%M%S')}-{secrets.token_hex(3)}"
        batch = PrintBatch(batch_id=batch_id, material_id=material.material_id, quantity=body.quantity, status="pending", printed_at=printed_at, created_by=user.id)
        labels = []
        label_payloads = []
        for _ in range(body.quantity):
            label_id = f"LBL-{printed_at.strftime('%Y%m%d')}-{secrets.token_hex(5)}"
            payload = {
                "v": 1,
                "labelId": label_id,
                "materialId": material.material_id,
                "materialCode": material.material_code,
                "name": material.name_zh,
                "printedAt": printed_at.isoformat(),
                "labelSize": label_size,
            }
            label_payloads.append(payload)
            labels.append(Label(label_id=label_id, material_id=material.material_id, payload_json=json.dumps(payload, ensure_ascii=False), print_batch_id=batch_id, printed_at=printed_at, created_by=user.id))
        db.add(batch)
        db.add_all(labels)
        db.flush()
        write_audit(
            db,
            actor_id=user.id,
            action="label.batch_created",
            resource_type="print_batch",
            resource_id=batch_id,
            detail={
                "material_id": material.material_id,
                "material_code": material.material_code,
                "quantity": body.quantity,
                "label_size": label_size,
            },
        )
        return {
            "batch_id": batch_id,
            "status": batch.status,
            "quantity": body.quantity,
            "printed_at": printed_at.isoformat(),
            "labels": [label.label_id for label in labels],
            "label_payloads": label_payloads,
            "label_size": label_size,
        }

    return execute_idempotent(
        db,
        request,
        user_id=user.id,
        endpoint="POST /labels/print-batches",
        payload={"material_id": body.material_id, "quantity": body.quantity},
        operation=operation,
        status_code=201,
    )


@router.get("/print-batches")
def list_print_batches(_: User = Depends(require_admin), db: Session = Depends(get_db)) -> list[dict]:
    batches = db.scalars(select(PrintBatch).order_by(PrintBatch.printed_at.desc())).all()
    materials = {item.material_id: item.name_zh for item in db.scalars(select(Material)).all()}
    users = {
        item.id: item.display_name or item.username
        for item in db.scalars(select(User).where(User.id.in_({batch.created_by for batch in batches}))).all()
    } if batches else {}
    return [
        {
            "batch_id": item.batch_id,
            "material_id": item.material_id,
            "material_name": materials.get(item.material_id, "已删除辅料"),
            "quantity": item.quantity,
            "status": item.status,
            "printed_at": item.printed_at.isoformat(),
            "created_by_name": users.get(item.created_by, "未知用户"),
        }
        for item in batches
    ]


@router.get("/excel-template")
def excel_template(_: User = Depends(require_admin)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "materials"
    sheet.append(MATERIAL_EXCEL_HEADERS)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=materials-template.xlsx"})


@router.get("/excel-export")
def excel_export(_: User = Depends(require_admin), db: Session = Depends(get_db)):
    materials = db.scalars(select(Material).order_by(Material.material_code)).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "materials"
    sheet.append(MATERIAL_EXCEL_HEADERS)
    for material in materials:
        sheet.append(
            (
                material.material_code,
                material.name_zh,
                material.name_en,
                material.shelf_life_months,
            )
        )
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 20
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=materials-export.xlsx"},
    )


@router.post("/excel-validate")
async def validate_excel(file: UploadFile = File(...), _: User = Depends(require_admin), db: Session = Depends(get_db)) -> dict:
    workbook = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or tuple(rows[0][:4]) != MATERIAL_EXCEL_HEADERS:
        raise HTTPException(status_code=422, detail={"code": "INVALID_TEMPLATE", "message": "Excel 模板表头不正确"})
    errors = []
    valid = []
    existing = {value for value in db.scalars(select(Material.material_code)).all()}
    for row_no, row in enumerate(rows[1:], start=2):
        code, name_zh, name_en, shelf = list(row)[:4] + [None] * max(0, 4 - len(row))
        if not code or not name_zh or not isinstance(shelf, (int, float)) or shelf < 0:
            errors.append({"row": row_no, "message": "代号、中文名称和非负保质期为必填且格式正确"})
        elif str(code) in existing or any(item["material_code"] == str(code) for item in valid):
            errors.append({"row": row_no, "message": "辅料代号已存在或在文件中重复"})
        else:
            valid.append({"material_code": str(code), "name_zh": str(name_zh), "name_en": str(name_en) if name_en else None, "shelf_life_months": int(shelf)})
    return {"valid_count": len(valid), "error_count": len(errors), "errors": errors, "valid_rows": valid}


@router.post("/excel-import")
async def import_excel(file: UploadFile = File(...), _: User = Depends(require_admin), db: Session = Depends(get_db)) -> dict:
    workbook = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or tuple(rows[0][:4]) != MATERIAL_EXCEL_HEADERS:
        raise HTTPException(status_code=422, detail={"code": "INVALID_TEMPLATE", "message": "Excel 模板表头不正确"})
    existing = {value for value in db.scalars(select(Material.material_code)).all()}
    created = []
    errors = []
    for row_no, row in enumerate(rows[1:], start=2):
        code, name_zh, name_en, shelf = list(row)[:4] + [None] * max(0, 4 - len(row))
        if not code or not name_zh or not isinstance(shelf, (int, float)) or shelf < 0:
            errors.append({"row": row_no, "message": "代号、中文名称和非负保质期为必填且格式正确"})
            continue
        material_code = str(code)
        if material_code in existing or any(item.material_code == material_code for item in created):
            errors.append({"row": row_no, "message": "辅料代号已存在或在文件中重复"})
            continue
        material = Material(
            material_id=f"MAT-{db.query(Material).count() + len(created) + 1:05d}",
            material_code=material_code,
            name_zh=str(name_zh),
            name_en=str(name_en) if name_en else None,
            shelf_life_months=int(shelf),
        )
        created.append(material)
    if created:
        db.add_all(created)
        db.commit()
    return {"imported_count": len(created), "error_count": len(errors), "errors": errors}
