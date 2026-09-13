from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


PRODUCT_RECIPE_HEADERS = ("product_name", "material_code", "quantity_per_ton_kg")
PRODUCT_EXPORT_HEADERS = (
    "product_name",
    "material_code",
    "quantity_per_ton_kg",
    "sort_order",
    "recipe_version",
    "product_enabled",
    "material_name",
)
MAX_IMPORT_ROWS = 5_000
MAX_IMPORT_PRODUCTS = 200
MAX_RECIPE_ITEMS = 100


class ProductExcelFormatError(ValueError):
    pass


@dataclass(frozen=True)
class ProductRecipeRow:
    row_no: int
    product_name: str
    material_code: str
    quantity_per_ton_kg: float


@dataclass(frozen=True)
class ParsedProductRecipe:
    product_name: str
    rows: tuple[ProductRecipeRow, ...]


@dataclass(frozen=True)
class ProductRecipeParseResult:
    valid_products: tuple[ParsedProductRecipe, ...]
    valid_rows: tuple[dict, ...]
    errors: tuple[dict, ...]


def _workbook_bytes(workbook: Workbook) -> BytesIO:
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_product_recipe_template() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品配方"
    sheet.append(PRODUCT_RECIPE_HEADERS)
    sheet.freeze_panes = "A2"
    for cell, hint in zip(
        sheet[1],
        (
            "产品名称，相同产品连续填写多行。",
            "必须与辅料管理中已有的辅料代号一致。",
            "每吨成品需要的辅料重量，单位 kg。",
        ),
    ):
        cell.comment = Comment(hint, "牧衡辅料称重防错系统")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 24
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287A5B")
        cell.alignment = Alignment(horizontal="center")

    guide = workbook.create_sheet("填写说明")
    guide.append(["项目", "说明"])
    guide.append(["必填表头", "product_name、material_code、quantity_per_ton_kg"])
    guide.append(["产品分组", "相同 product_name 的多行会组成一个产品配方，行的先后顺序就是步骤顺序。"])
    guide.append(["辅料", "material_code 必须已存在于辅料管理中，系统不会自动创建辅料。"])
    guide.append(["每吨用量", "必须大于 0，单位为 kg/吨；示例 1.25 表示每吨成品使用 1.25 kg。"])
    guide.append(["导入规则", "任意一行校验失败时，该产品的所有行都不会写入；已有产品不会被覆盖。"])
    guide[1][0].font = Font(bold=True)
    guide[1][1].font = Font(bold=True)
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 90
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return _workbook_bytes(workbook)


def build_product_recipe_export(products: Iterable) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品配方"
    sheet.append(PRODUCT_EXPORT_HEADERS)
    for product in products:
        recipe = product.recipe
        for item in sorted(recipe.items, key=lambda value: value.sort_order) if recipe else []:
            sheet.append(
                (
                    product.name,
                    item.material.material_code,
                    item.quantity_per_ton_kg,
                    item.sort_order + 1,
                    recipe.version,
                    "是" if product.enabled else "否",
                    item.material.name_zh,
                )
            )
    sheet.freeze_panes = "A2"
    widths = (28, 18, 24, 12, 14, 16, 22)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="287A5B")
        cell.alignment = Alignment(horizontal="center")
    return _workbook_bytes(workbook)


def parse_product_recipe_workbook(
    content: bytes,
    *,
    existing_material_codes: set[str],
    existing_product_names: set[str],
) -> ProductRecipeParseResult:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ProductExcelFormatError("Excel 文件无法读取，请确认文件为有效的 .xlsx 工作簿") from exc

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ProductExcelFormatError("Excel 文件没有数据")

    header_row_index = None
    for index, row in enumerate(rows[:10]):
        values = tuple(row[: len(PRODUCT_RECIPE_HEADERS)])
        if values == PRODUCT_RECIPE_HEADERS:
            header_row_index = index
            break
    if header_row_index is None:
        raise ProductExcelFormatError(
            "Excel 表头不正确，前三列必须依次为 product_name、material_code、quantity_per_ton_kg"
        )

    data_rows = rows[header_row_index + 1 :]
    non_empty_count = sum(1 for row in data_rows if any(value is not None and str(value).strip() for value in row))
    if non_empty_count > MAX_IMPORT_ROWS:
        raise ProductExcelFormatError(f"单次最多导入 {MAX_IMPORT_ROWS} 行配方数据")

    errors: list[dict] = []
    error_rows: set[int] = set()
    product_names_with_errors: set[str] = set()
    grouped: dict[str, list[ProductRecipeRow]] = {}
    for offset, row in enumerate(data_rows, start=header_row_index + 2):
        values = list(row[: len(PRODUCT_RECIPE_HEADERS)]) + [None] * max(
            0, len(PRODUCT_RECIPE_HEADERS) - len(row)
        )
        if not any(value is not None and str(value).strip() for value in values):
            continue

        product_name = str(values[0]).strip() if values[0] is not None else ""
        material_code = str(values[1]).strip() if values[1] is not None else ""
        quantity = values[2]

        row_error = False
        if not product_name or len(product_name) > 128:
            errors.append({"row": offset, "product_name": product_name, "material_code": material_code, "message": "产品名称为必填项且不能超过 128 字"})
            row_error = True
        if not material_code or len(material_code) > 64:
            errors.append({"row": offset, "product_name": product_name, "material_code": material_code, "message": "辅料代号为必填项且不能超过 64 字"})
            row_error = True
        elif material_code not in existing_material_codes:
            errors.append({"row": offset, "product_name": product_name, "material_code": material_code, "message": "辅料代号不存在，请先在辅料管理中创建"})
            row_error = True
        if isinstance(quantity, bool) or not isinstance(quantity, (int, float)) or quantity <= 0 or quantity > 10_000:
            errors.append({"row": offset, "product_name": product_name, "material_code": material_code, "message": "每吨用量必须为大于 0 且不超过 10000 的数字"})
            row_error = True

        if row_error:
            error_rows.add(offset)
            if product_name:
                product_names_with_errors.add(product_name)
            continue
        grouped.setdefault(product_name, []).append(
            ProductRecipeRow(
                row_no=offset,
                product_name=product_name,
                material_code=material_code,
                quantity_per_ton_kg=float(quantity),
            )
        )

    if len(grouped) > MAX_IMPORT_PRODUCTS:
        raise ProductExcelFormatError(f"单次最多导入 {MAX_IMPORT_PRODUCTS} 个产品")

    for product_name, product_rows in grouped.items():
        group_has_error = product_name in product_names_with_errors
        if product_name in existing_product_names:
            for row in product_rows:
                if row.row_no not in error_rows:
                    errors.append({"row": row.row_no, "product_name": product_name, "material_code": row.material_code, "message": "产品名称已存在；为避免覆盖，本产品不会导入"})
                    error_rows.add(row.row_no)
            group_has_error = True

        seen_materials: set[str] = set()
        for row in product_rows:
            if row.material_code in seen_materials:
                if row.row_no not in error_rows:
                    errors.append({"row": row.row_no, "product_name": product_name, "material_code": row.material_code, "message": "同一产品中辅料不能重复"})
                    error_rows.add(row.row_no)
                group_has_error = True
            seen_materials.add(row.material_code)

        if len(product_rows) > MAX_RECIPE_ITEMS:
            for row in product_rows:
                if row.row_no not in error_rows:
                    errors.append({"row": row.row_no, "product_name": product_name, "material_code": row.material_code, "message": f"单个产品最多包含 {MAX_RECIPE_ITEMS} 种辅料"})
                    error_rows.add(row.row_no)
            group_has_error = True

        if group_has_error:
            grouped[product_name] = []

    valid_products = tuple(
        ParsedProductRecipe(product_name=name, rows=tuple(rows))
        for name, rows in grouped.items()
        if rows
    )
    valid_rows = tuple(
        {
            "row": row.row_no,
            "product_name": row.product_name,
            "material_code": row.material_code,
            "quantity_per_ton_kg": row.quantity_per_ton_kg,
        }
        for product in valid_products
        for row in product.rows
    )
    errors.sort(key=lambda item: item["row"])
    return ProductRecipeParseResult(
        valid_products=valid_products,
        valid_rows=valid_rows,
        errors=tuple(errors),
    )
