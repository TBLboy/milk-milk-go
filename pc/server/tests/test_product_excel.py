from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.product_excel import (
    PRODUCT_EXPORT_HEADERS,
    PRODUCT_RECIPE_HEADERS,
)


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def admin_headers(client):
    token = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "admin123"},
    ).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def create_material(client, headers, *, code: str, name: str) -> dict:
    response = client.post(
        "/api/v1/master-data/materials",
        headers=headers,
        json={
            "material_code": code,
            "name_zh": name,
            "shelf_life_months": 12,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def recipe_workbook(rows: list[tuple]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品配方"
    sheet.append(PRODUCT_RECIPE_HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_product_recipe_template_has_expected_headers_and_guide(client):
    response = client.get(
        "/api/v1/master-data/products/import-template",
        headers=admin_headers(client),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(EXCEL_MIME)
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["产品配方", "填写说明"]
    recipe_sheet = workbook["产品配方"]
    assert tuple(cell.value for cell in recipe_sheet[1]) == PRODUCT_RECIPE_HEADERS
    assert "product_name、material_code、quantity_per_ton_kg" in workbook["填写说明"]["B2"].value
    assert "不会自动创建辅料" in workbook["填写说明"]["B4"].value


def test_product_recipe_export_uses_current_version_only(client):
    headers = admin_headers(client)
    material_a = create_material(client, headers, code="EXP-A1", name="导出辅料 A")
    material_b = create_material(client, headers, code="EXP-B2", name="导出辅料 B")
    created = client.post(
        "/api/v1/master-data/products",
        headers=headers,
        json={
            "name": "导出测试产品",
            "items": [
                {"material_id": material_a["material_id"], "quantity_per_ton_kg": 4},
                {"material_id": material_b["material_id"], "quantity_per_ton_kg": 3},
            ],
        },
    )
    assert created.status_code == 201, created.text
    product_id = created.json()["id"]
    updated = client.put(
        f"/api/v1/master-data/products/{product_id}",
        headers=headers,
        json={
            "name": "导出测试产品",
            "items": [
                {"material_id": material_b["material_id"], "quantity_per_ton_kg": 2.5},
                {"material_id": material_a["material_id"], "quantity_per_ton_kg": 7.5},
            ],
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["recipe_version"] == 2

    response = client.get("/api/v1/master-data/products/export", headers=headers)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook.active
    assert tuple(cell.value for cell in sheet[1]) == PRODUCT_EXPORT_HEADERS
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows == [
        ("导出测试产品", "EXP-B2", 2.5, 1, 2, "是", "导出辅料 B"),
        ("导出测试产品", "EXP-A1", 7.5, 2, 2, "是", "导出辅料 A"),
    ]
    assert all(row[2] != 4.0 for row in rows)


def test_validate_and_import_return_group_errors_without_overwrite(client):
    headers = admin_headers(client)
    material_a = create_material(client, headers, code="ATOM-A1", name="原子导入辅料 A")
    material_b = create_material(client, headers, code="ATOM-B2", name="原子导入辅料 B")
    existing = client.post(
        "/api/v1/master-data/products",
        headers=headers,
        json={
            "name": "已有配方",
            "items": [
                {"material_id": material_a["material_id"], "quantity_per_ton_kg": 4},
            ],
        },
    )
    assert existing.status_code == 201, existing.text

    content = recipe_workbook(
        [
            ("有效产品", "ATOM-A1", 5),
            ("有效产品", "ATOM-B2", 2),
            ("重复辅料产品", "ATOM-A1", 1),
            ("重复辅料产品", "ATOM-A1", 2),
            ("未知辅料产品", "NOT-EXIST", 1),
            ("未知辅料产品", "ATOM-B2", 3),
            ("非法用量产品", "ATOM-A1", 0),
            ("已有配方", "ATOM-A1", 9),
        ]
    )
    files = {"file": ("recipes.xlsx", content, EXCEL_MIME)}
    validated = client.post(
        "/api/v1/master-data/products/import-validate",
        headers=headers,
        files=files,
    )
    assert validated.status_code == 200, validated.text
    validation = validated.json()
    assert validation["valid_product_count"] == 1
    assert validation["valid_row_count"] == 2
    assert [item["product_name"] for item in validation["valid_products"]] == ["有效产品"]
    messages = [item["message"] for item in validation["errors"]]
    assert "辅料代号不存在，请先在辅料管理中创建" in messages
    assert "同一产品中辅料不能重复" in messages
    assert "每吨用量必须为大于 0 且不超过 10000 的数字" in messages
    assert "产品名称已存在；为避免覆盖，本产品不会导入" in messages

    imported = client.post(
        "/api/v1/master-data/products/import",
        headers=headers,
        files={"file": ("recipes.xlsx", content, EXCEL_MIME)},
    )
    assert imported.status_code == 200, imported.text
    result = imported.json()
    assert result["imported_product_count"] == 1
    assert result["imported_row_count"] == 2
    assert result["imported_products"] == ["有效产品"]
    assert result["error_count"] == len(validation["errors"])

    products = {
        item["name"]: item
        for item in client.get("/api/v1/master-data/products", headers=headers).json()
    }
    assert set(products) == {"已有配方", "有效产品"}
    assert [item["material_code"] for item in products["有效产品"]["items"]] == [
        "ATOM-A1",
        "ATOM-B2",
    ]
    assert products["有效产品"]["recipe_version"] == 1
    assert products["已有配方"]["recipe_version"] == 1
    assert products["已有配方"]["items"][0]["quantity_per_ton_kg"] == 4

    created_order = client.post(
        "/api/v1/work-orders",
        headers=headers,
        json={"product_id": products["有效产品"]["id"], "target_weight_kg": 500},
    )
    assert created_order.status_code == 201, created_order.text
    assert [step["required_weight_kg"] for step in created_order.json()["steps"]] == [
        2.5,
        1.0,
    ]

    audit_logs = client.get(
        "/api/v1/operations/audit-logs",
        headers=headers,
        params={"action": "product.excel_imported"},
    ).json()
    assert audit_logs["total"] == 1
    detail = audit_logs["items"][0]["detail"]
    assert detail["imported_products"] == ["有效产品"]
    assert detail["imported_row_count"] == 2


def test_product_recipe_excel_endpoints_reject_operator(client):
    headers = admin_headers(client)
    created = client.post(
        "/api/v1/auth/users",
        headers=headers,
        json={
            "username": "recipe_operator",
            "display_name": "配方操作员",
            "password": "operator123",
            "employee_no": "MH-EXCEL-01",
        },
    )
    assert created.status_code in {200, 201}, created.text
    token = client.post(
        "/api/v1/auth/login",
        json={"username": "recipe_operator", "password": "operator123"},
    ).json()["access_token"]
    operator = {"Authorization": f"Bearer {token}"}
    content = recipe_workbook([("无权产品", "ANY", 1)])

    assert client.get(
        "/api/v1/master-data/products/import-template",
        headers=operator,
    ).status_code == 403
    assert client.get(
        "/api/v1/master-data/products/export",
        headers=operator,
    ).status_code == 403
    assert client.post(
        "/api/v1/master-data/products/import-validate",
        headers=operator,
        files={"file": ("recipes.xlsx", content, EXCEL_MIME)},
    ).status_code == 403
    assert client.post(
        "/api/v1/master-data/products/import",
        headers=operator,
        files={"file": ("recipes.xlsx", content, EXCEL_MIME)},
    ).status_code == 403
