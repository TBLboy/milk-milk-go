from io import BytesIO

from openpyxl import Workbook, load_workbook


def admin_headers(client):
    token = client.post("/api/v1/auth/login", json={"username": "admin", "password": "admin123"}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_admin_can_list_created_users(client):
    headers = admin_headers(client)
    create = client.post(
        "/api/v1/auth/users",
        headers=headers,
        json={"username": "operator99", "display_name": "测试操作员", "password": "password123", "employee_no": "MH1005"},
    )
    assert create.status_code == 201
    users = client.get("/api/v1/auth/users", headers=headers).json()
    assert any(item["username"] == "operator99" for item in users)


def test_admin_can_reset_password_and_toggle_operator_active(client):
    headers = admin_headers(client)
    user = client.post(
        "/api/v1/auth/users",
        headers=headers,
        json={"username": "operator98", "display_name": "重置测试", "password": "password123", "employee_no": "MH1006"},
    ).json()["user"]
    reset = client.post(f"/api/v1/auth/users/{user['id']}/reset-password", headers=headers)
    assert reset.status_code == 200
    reset_password = reset.json()["temporary_password"]
    login = client.post("/api/v1/auth/login", json={"username": "operator98", "password": reset_password})
    assert login.status_code == 200
    deactivate = client.patch(f"/api/v1/auth/users/{user['id']}/active", headers=headers, json={"is_active": False})
    assert deactivate.status_code == 200
    assert deactivate.json()["user"]["is_active"] is False
    assert client.post("/api/v1/auth/login", json={"username": "operator98", "password": reset_password}).status_code == 403


def test_settings_defaults_can_be_updated(client):
    headers = admin_headers(client)
    defaults = client.get("/api/v1/settings", headers=headers)
    assert defaults.status_code == 200
    assert defaults.json()["default_tolerance_percent"] == "1.0"
    updated = client.put("/api/v1/settings", headers=headers, json={"values": {"default_tolerance_percent": "1.5"}})
    assert updated.status_code == 200
    assert updated.json()["default_tolerance_percent"] == "1.5"


def test_excel_import_creates_materials(client):
    headers = admin_headers(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "materials"
    sheet.append(["material_code", "name_zh", "name_en", "shelf_life_months"])
    sheet.append(["Z9", "测试辅料", "Test Material", 18])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = client.post("/api/v1/labels/excel-import", headers=headers, files={"file": ("materials.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    assert response.json()["imported_count"] == 1
    materials = client.get("/api/v1/master-data/materials", headers=headers).json()
    assert any(item["material_code"] == "Z9" for item in materials)


def test_excel_export_contains_all_materials_in_code_order(client):
    headers = admin_headers(client)
    for code, name_zh, name_en, shelf_life in (
        ("Z9", "测试辅料 Z", "Test Material Z", 18),
        ("A9", "测试辅料 A", "Test Material A", 24),
    ):
        response = client.post(
            "/api/v1/master-data/materials",
            headers=headers,
            json={
                "material_code": code,
                "name_zh": name_zh,
                "name_en": name_en,
                "shelf_life_months": shelf_life,
            },
        )
        assert response.status_code == 201

    response = client.get("/api/v1/labels/excel-export", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook.active
    assert tuple(cell.value for cell in sheet[1]) == (
        "material_code",
        "name_zh",
        "name_en",
        "shelf_life_months",
    )
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows == [
        ("A9", "测试辅料 A", "Test Material A", 24),
        ("Z9", "测试辅料 Z", "Test Material Z", 18),
    ]


def test_excel_export_empty_library_returns_valid_workbook(client):
    response = client.get("/api/v1/labels/excel-export", headers=admin_headers(client))
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook.active
    assert tuple(cell.value for cell in sheet[1]) == (
        "material_code",
        "name_zh",
        "name_en",
        "shelf_life_months",
    )
    assert sheet.max_row == 1


def test_excel_export_rejects_operator(client):
    headers = admin_headers(client)
    created = client.post(
        "/api/v1/auth/users",
        headers=headers,
        json={
            "username": "operator93",
            "display_name": "导出权限测试",
            "password": "password123",
            "employee_no": "MH1093",
        },
    )
    assert created.status_code == 201
    login = client.post(
        "/api/v1/auth/login",
        json={"username": "operator93", "password": "password123"},
    )
    assert login.status_code == 200
    operator_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    response = client.get("/api/v1/labels/excel-export", headers=operator_headers)
    assert response.status_code == 403
